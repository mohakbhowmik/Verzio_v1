import os
from datetime import datetime, date, time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session
from database import Appointment

def generate_excel_report(business_id: int, business_name: str, db: Session, from_date: date, to_date: date) -> str:
    # 1. Fetch data for the custom range
    start_dt = datetime.combine(from_date, time.min)
    end_dt = datetime.combine(to_date, time.max)
    
    appointments = db.query(Appointment).filter(
        Appointment.business_id == business_id,
        Appointment.appointment_time >= start_dt,
        Appointment.appointment_time <= end_dt
    ).order_by(Appointment.appointment_time.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Business Report"

    # --- Style Definitions ---
    header_navy = "1E293B"
    sub_header_gray = "64748B"
    border_color = "E2E8F0"
    zebra_fill = "F8FAFC"
    white = "FFFFFF"

    title_font = Font(name='Segoe UI', size=16, bold=True, color=header_navy)
    meta_font = Font(name='Segoe UI', size=10, color=sub_header_gray)
    col_header_font = Font(name='Segoe UI', size=11, bold=True, color=white)
    data_font = Font(name='Segoe UI', size=11, color="334155")
    
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", indent=1)
    thin_border = Border(bottom=Side(style='thin', color=border_color))

    # --- Header Section ---
    ws.merge_cells('A1:E1')
    ws['A1'] = business_name.upper()
    ws['A1'].font = title_font

    ws.merge_cells('A2:E2')
    range_str = f"REPORT PERIOD: {from_date.strftime('%d %b %Y')} to {to_date.strftime('%d %b %Y')}"
    ws['A2'] = range_str
    ws['A2'].font = meta_font
    
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 18

    # --- Table Column Headers (Added DATE) ---
    headers = ["DATE", "TIME", "CUSTOMER NAME", "PHONE", "STATUS"]
    header_fill = PatternFill(start_color=header_navy, end_color=header_navy, fill_type="solid")
    
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header_title
        cell.font = col_header_font
        cell.fill = header_fill
        cell.alignment = center_align
    
    ws.row_dimensions[4].height = 25

    # --- Populate Data ---
    current_row = 5
    for appt in appointments:
        row_data = [
            appt.appointment_time.strftime("%d-%b-%Y"), # Date
            appt.appointment_time.strftime("%I:%M %p"), # Time
            (appt.customer_name or "Guest").title(),
            appt.customer_phone,
            appt.status.upper()
        ]
        
        fill = PatternFill(start_color=zebra_fill, end_color=zebra_fill, fill_type="solid") if current_row % 2 == 0 else None

        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_num, value=value)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align if col_num != 3 else left_align
            if fill:
                cell.fill = fill
            
            if col_num == 5: # Status coloring
                colors = {"CONFIRMED": "16A34A", "COMPLETED": "2563EB", "CANCELLED": "DC2626", "PENDING": "D97706"}
                cell.font = Font(name='Segoe UI', size=10, bold=True, color=colors.get(value, "64748B"))

        ws.row_dimensions[current_row].height = 22
        current_row += 1

    # --- Column Widths ---
    widths = [15, 12, 25, 18, 15]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = width

    ws.views.sheetView[0].showGridLines = False

    # Save with range-specific filename
    os.makedirs("exports", exist_ok=True)
    fname = f"exports/Report_{business_id}_{from_date.isoformat()}_to_{to_date.isoformat()}.xlsx"
    wb.save(fname)
    return fname